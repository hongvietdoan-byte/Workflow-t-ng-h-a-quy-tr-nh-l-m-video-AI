import io
import os
import tempfile
import unicodedata
import unittest
import zipfile

import openpyxl
from streamlit.testing.v1 import AppTest

from core import script_parser, script_reader
from core.script_reader import ScriptReadError, from_text, read_script
from core.db import connect
from core.pipeline import Pipeline

APP = os.path.join(os.path.dirname(__file__), "..", "dashboard", "app.py")
HEADER = ["Cảnh", "Bối cảnh", "Thời gian", "Mô tả", "Nhân vật", "Lời thoại", "Góc máy"]
ROWS = [HEADER,
        ["1", "Rừng Elder", "Đêm", "Sương mù phủ kín khu rừng.", "Lyra", "Có thứ gì đó đang theo chúng ta.", "Cận cảnh"],
        ["", "", "", "", "Kael", "Ở yên sau lưng ta.", ""],                                  # same scene, another line
        ["2", "Pháo đài", "Ngày", "Ông lão Orin nhìn xuống thung lũng.", "Orin", "Bóng Đêm đã tỉnh giấc.", ""]]


def docx_with(body_xml: str) -> bytes:
    xml = ('<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f"<w:body>{body_xml}</w:body></w:document>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def p(text):
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def table(rows):
    cells = lambda r: "".join(f"<w:tc>{p(c)}</w:tc>" for c in r)  # noqa: E731
    return "<w:tbl>" + "".join(f"<w:tr>{cells(r)}</w:tr>" for r in rows) + "</w:tbl>"


def xlsx_bytes(sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows, *hidden in sheets:
        ws = book.create_sheet(name)
        for r in rows:
            ws.append(r)
        if hidden:
            ws.sheet_state = "hidden"
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def scenes_of(res):
    return script_parser.split_scenes(res.paragraphs)


class TableScriptTests(unittest.TestCase):
    def check_table_result(self, res):
        scenes = scenes_of(res)
        self.assertEqual(len(scenes), 2)
        self.assertEqual(scenes[0].heading, "CẢNH 1 - Đêm, Rừng Elder")
        self.assertIn("Sương mù phủ kín khu rừng.", scenes[0].text)
        self.assertIn("LYRA: Có thứ gì đó đang theo chúng ta.", scenes[0].text)
        self.assertIn("KAEL: Ở yên sau lưng ta.", scenes[0].text)                  # the blank-scene row continued scene 1
        self.assertIn("(Góc máy: Cận cảnh)", scenes[0].text)
        self.assertEqual(scenes[0].characters, ["KAEL", "LYRA"])                     # notes are not mistaken for speakers
        self.assertEqual(scenes[1].heading, "CẢNH 2 - Ngày, Pháo đài")

    def test_word_table(self):
        self.check_table_result(read_script("kich_ban.docx", docx_with(p("Trailer Ep.1") + table(ROWS))))

    def test_excel_sheet(self):
        self.check_table_result(read_script("kich_ban.xlsx", xlsx_bytes([("Kịch bản", ROWS)])))

    def test_csv_with_semicolons_and_a_vietnamese_byte_order_mark(self):
        text = "\n".join(";".join(r) for r in ROWS)
        self.check_table_result(read_script("kb.csv", text.encode("utf-8-sig")))

    def test_pasted_block_copied_from_excel(self):
        text = "\n".join("\t".join(r) for r in ROWS)
        res = from_text(text)
        self.check_table_result(res)
        self.assertIn("Excel", res.info[0])

    def test_english_headers_and_scene_numbers_are_understood(self):
        rows = [["Scene", "Description", "Dialogue", "Character", "Duration"],
                ["1", "Hero enters the arena.", "I am ready.", "Hero", "5"],
                ["2", "The crowd cheers.", "", "", "4"]]
        res = read_script("x.xlsx", xlsx_bytes([("Sheet1", rows)]))
        scenes = scenes_of(res)
        self.assertEqual(len(scenes), 2)
        self.assertIn("HERO: I am ready.", scenes[0].text)
        self.assertIn("(Thời lượng: 5)", scenes[0].text)
        self.assertTrue(any("Description → description" in line for line in res.info))

    def test_speaker_missing_gets_a_placeholder_and_inline_names_are_kept(self):
        rows = [["Cảnh", "Mô tả", "Thoại"], ["1", "Hai người gặp nhau.", "A: Chào.\nB: Chào lại."], ["2", "Một mình.", "Đi thôi."]]
        scenes = scenes_of(read_script("x.csv", "\n".join(",".join(f'"{c}"' for c in r) for r in rows).encode("utf-8")))
        self.assertIn("A: Chào.", scenes[0].text)
        self.assertIn("B: Chào lại.", scenes[0].text)
        self.assertIn("LỜI THOẠI: Đi thôi.", scenes[1].text)

    def test_without_a_scene_column_a_new_description_starts_a_scene(self):
        rows = [["Mô tả", "Nhân vật", "Lời thoại"], ["Cửa mở.", "Lyra", "Ai đó?"], ["", "Kael", "Là ta."], ["Trời sáng.", "", ""]]
        scenes = scenes_of(read_script("x.xlsx", xlsx_bytes([("S", rows)])))
        self.assertEqual(len(scenes), 2)
        self.assertIn("KAEL: Là ta.", scenes[0].text)

    def test_a_table_without_a_header_becomes_one_scene_per_row_and_says_so(self):
        res = read_script("x.csv", "Lyra bước vào rừng,đêm\nKael rút kiếm,đêm".encode("utf-8"))
        self.assertEqual(len(scenes_of(res)), 2)
        self.assertTrue(any("không có dòng tiêu đề" in line for line in res.info))

    def test_only_script_sheets_of_a_workbook_are_used_and_scene_numbers_continue(self):
        book = xlsx_bytes([("Ghi chú", [["Việc cần làm"], ["mua cà phê"]]), ("Tập 1", ROWS), ("Tập 2", ROWS[:2]),
                           ("Ẩn", ROWS, True)])
        res = read_script("x.xlsx", book)
        headings = [s.heading for s in scenes_of(res)]
        self.assertEqual([h.split(" - ")[0] for h in headings], ["CẢNH 1", "CẢNH 2", "CẢNH 3"])
        self.assertTrue(any("Bỏ qua sheet" in line and "Ghi chú" in line for line in res.info))

    def test_word_file_with_normal_paragraphs_and_a_table_keeps_the_order(self):
        body = p("CẢNH 1 - ĐÊM") + p("Mở đầu bằng chữ.") + p("LYRA: Xin chào.") + table(ROWS)
        scenes = scenes_of(read_script("mixed.docx", docx_with(body)))
        self.assertEqual(len(scenes), 3)
        self.assertEqual(scenes[0].heading, "CẢNH 1 - ĐÊM")
        self.assertEqual(scenes[1].heading.split(" - ")[0], "CẢNH 1")                 # numbering of the table starts at 1 again


class PlainTextTests(unittest.TestCase):
    def test_typed_text_is_split_into_lines_and_scenes(self):
        res = from_text("CẢNH 1 - ĐÊM\nLyra chạy.\n\nLYRA: Nhanh lên!\nCẢNH 2\nSáng rồi.")
        scenes = scenes_of(res)
        self.assertEqual([s.heading for s in scenes], ["CẢNH 1 - ĐÊM", "CẢNH 2"])
        self.assertEqual(scenes[0].characters, ["LYRA"])

    def test_utf8_utf16_and_old_windows_encodings_are_decoded(self):
        line = "CẢNH 1\nLYRA: Tiếng Việt có dấu."
        for enc in ("utf-8", "utf-8-sig", "utf-16"):
            self.assertIn("Tiếng Việt có dấu", read_script("k.txt", line.encode(enc)).paragraphs[-1], enc)
        self.assertIn("Café", read_script("k.txt", "CANH 1\nLYRA: Café".encode("cp1252")).paragraphs[-1])   # legacy Windows text

    def test_a_pasted_script_with_tabs_only_in_a_few_lines_is_still_plain_text(self):
        res = from_text("CẢNH 1\nLyra\tchạy nhanh\nSáng rồi.\nKael đứng lại.")
        self.assertEqual(len(res.paragraphs), 4)


class RefusalTests(unittest.TestCase):
    def test_clear_messages_for_bad_or_unsupported_input(self):
        for name, data, word in (("a.pdf", b"%PDF", "PDF"), ("a.xls", b"x", "Lưu thành"), ("a.exe", b"x", "Chưa hỗ trợ"),
                                 ("a.docx", b"not a zip", "docx"), ("a.xlsx", b"not a zip", "xlsx"), ("a.txt", b"", "rỗng")):
            with self.assertRaises(ScriptReadError) as ctx:
                read_script(name, data)
            self.assertIn(word, str(ctx.exception), name)
        with self.assertRaises(ScriptReadError):
            from_text("   ")
        with self.assertRaises(ScriptReadError):
            read_script("none.xlsx", xlsx_bytes([("A", [["x", "y"], ["1", "2"]]), ("B", [["q"], ["w"]])]))   # no sheet looks like a script
        with self.assertRaises(ScriptReadError):
            read_script("empty.docx", docx_with(""))


class DashboardTests(unittest.TestCase):
    def test_typing_a_script_and_pressing_analyse_creates_the_scenes(self):
        tmp = tempfile.mkdtemp()
        os.environ.update({"PIPELINE_DB": os.path.join(tmp, "m.sqlite"), "PIPELINE_DATA": os.path.join(tmp, "projects")})
        try:
            pl = Pipeline(connect(os.environ["PIPELINE_DB"]))
            pid = pl.create_project("typed")
            at = AppTest.from_file(APP, default_timeout=40).run()
            self.assertFalse(at.exception)
            self.assertIn("📎 Tải file", [t.label for t in at.tabs])
            self.assertIn("✍ Gõ / dán văn bản", [t.label for t in at.tabs])
            at.text_area(key=f"paste_{pid}").set_value("CẢNH 1 - ĐÊM\nLyra chạy.\nLYRA: Nhanh lên!\nCẢNH 2\nSáng rồi.").run()
            next(b for b in at.button if b.key == f"btn_analyse_{pid}").click().run()
            self.assertFalse(at.exception)
            rows = pl.conn.execute("SELECT idx FROM scenes WHERE project_id=? ORDER BY idx", (pid,)).fetchall()
            self.assertEqual([r["idx"] for r in rows], [1, 2])
            self.assertTrue(any("Hệ thống đã đọc kịch bản" in e.label for e in at.expander))
        finally:
            os.environ.pop("PIPELINE_DB", None)
            os.environ.pop("PIPELINE_DATA", None)


if __name__ == "__main__":
    unittest.main()
